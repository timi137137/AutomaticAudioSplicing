"""Read the download schedule from an Excel workbook."""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from aas.domain.schedule import SheetEntry, parse_time_range
from aas.errors import SheetError

COL_CODE = "编码"
COL_TITLE = "歌名"
COL_RANGE = "时间区间"
COL_URL = "音源地址"
REQUIRED_HEADERS = (COL_CODE, COL_TITLE, COL_RANGE, COL_URL)


def _locate_columns(header: Iterable[Any]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, value in enumerate(header):
        name = str(value).strip() if value is not None else ""
        if name in REQUIRED_HEADERS and name not in columns:
            columns[name] = index
    missing = [name for name in REQUIRED_HEADERS if name not in columns]
    if missing:
        raise SheetError(f"Excel 缺少必需列: {'、'.join(missing)}")
    return columns


def _cell_text(row: tuple[Any, ...], index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _read_entries(
    rows: Iterator[tuple[Any, ...]], columns: dict[str, int]
) -> list[SheetEntry]:
    entries: list[SheetEntry] = []
    seen: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        code = _cell_text(row, columns[COL_CODE])
        title = _cell_text(row, columns[COL_TITLE])
        range_text = _cell_text(row, columns[COL_RANGE])
        url = _cell_text(row, columns[COL_URL])
        if not code:
            raise SheetError(f"第 {row_number} 行缺少编码。")
        if not title:
            raise SheetError(f"第 {row_number} 行缺少歌名。")
        if not range_text:
            raise SheetError(f"第 {row_number} 行缺少时间区间。")
        if not url:
            raise SheetError(f"第 {row_number} 行缺少音源地址。")
        if code in seen:
            raise SheetError(f"第 {row_number} 行编码重复: {code}")
        seen.add(code)
        try:
            start, end = parse_time_range(range_text)
        except SheetError as exc:
            raise SheetError(f"第 {row_number} 行: {exc}") from exc
        entries.append(
            SheetEntry(
                code=code,
                title=title,
                start_seconds=start,
                end_seconds=end,
                url=url,
            )
        )
    return entries


def read_sheet(path: Path) -> list[SheetEntry]:
    if not path.is_file():
        raise SheetError(f"Excel 文件不存在: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise SheetError("Excel 没有工作表。")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise SheetError("Excel 为空。")
        columns = _locate_columns(header)
        entries = _read_entries(rows, columns)
    finally:
        workbook.close()
    if not entries:
        raise SheetError("Excel 没有有效数据行。")
    return entries
